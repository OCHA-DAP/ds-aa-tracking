"""Sync the media manifest (site_src/media.json) from the folder "AA visuals", and keep the
folder on the dev blob.

The Media page (scripts/page_media.py) shows what the manifest lists, and nothing else.

The folder lives in the OCHA-HFSA SharePoint, "(08) External communication/AA visuals" (a
private group: nothing there can be read without signing in). It was a shared Google Drive
folder until that was deleted in October 2026; the SharePoint copy is the same files, name
for name and byte for byte. Since nothing can list SharePoint without a sign-in, a sync
starts from a copy of the folder on disk: open the folder in SharePoint, "Download" (one
zip of the whole folder), unzip it, and point this script at the unzipped "AA visuals".

    uv run python scripts/sync_media.py --from "DIR/AA visuals" --dry-run   # what would change
    uv run python scripts/sync_media.py --from "DIR/AA visuals"             # do it
    uv run python scripts/sync_media.py --fetch site_media                  # the publish step

A sync does three things:
- archives every file of the folder on the dev blob, path for path, under ORIGINALS, with
  an index (path, size, sha256): the originals (6 GB on 2026-10-27) are kept there so that
  losing the folder once more loses nothing. A file already there with the same size is not
  sent again;
- makes two renditions of each photo (WIDTHS: the grid's and the viewer's), turned upright
  from the camera's orientation, and keeps them on the dev blob under WEB, named by the
  photo's content (the first 16 hex digits of its sha256): renditions are made once, here,
  never during the build, so a photo's rendition never changes under the same name;
- rewrites the manifest: each photo and document with its key, name and size, and the
  social posts of the catalogue kept at the folder's root (its posts sheet).

--fetch is the other half, for scripts/publish.sh: it fetches the renditions the manifest
lists (and only those) from the blob into a folder the publish copies to gh-pages under
media/. A rendition the blob does not have stops the publish.

Run by hand, and read the diff before committing it: what is in the manifest is published
(the publish workflow runs on a push that touches site_src/). It is not part of the nightly
build on purpose, so nothing reaches the page that nobody looked at. The renditions on
gh-pages are public files (gh-pages is a public branch; the page's encryption covers its
HTML only), as the Drive folder's own renditions were when it was shared by link.

What is left out of the manifest, and said so on every run (all of it is still archived):
- the catalogue itself and any other spreadsheet at the folder's root (working files, with
  research notes), and hidden sidecar files (.BridgeSort, …);
- catalogue rows that are not a post on a public platform: rows with no exact post link
  (still in the catalogue's recovery queue), and leads that point to a file share or a
  production board (Dropbox, Trello) — private links are not ours to publish;
- of each post, everything but the post itself: the catalogue's notes, verification trail
  and local paths stay in the catalogue.

The manifest is generated: hand edits are overwritten by the next run. Curation (hiding an
item, a caption, a credit) will need its own file or table; none exists yet.
"""

import argparse
import hashlib
import io
import json
import mimetypes
import re
import sys
import urllib.parse
from datetime import UTC, date, datetime
from pathlib import Path

SHAREPOINT = {"site": "https://unitednations.sharepoint.com",
              "library": "/sites/OCHAHFSA/Shared Documents",
              "path": "/sites/OCHAHFSA/Shared Documents/(08) External communication/AA visuals"}
CATALOGUE_PREFIX = "AA Social Media & Visual Asset Catalogue"
POSTS_SHEET = "Verified posts"
MANIFEST = Path(__file__).parents[1] / "site_src" / "media.json"

CONTAINER = "projects"                                    # the dev blob
ORIGINALS = "ds-aa-tracking/raw/media/aa_visuals"         # the folder, path for path
INDEX = "ds-aa-tracking/raw/media/aa_visuals_index.json"  # path, bytes, sha256 of each file
WEB = "ds-aa-tracking/processed/media/web"                # <key>-w<width>.jpg
WIDTHS = (400, 1600)                                      # the grid's, the viewer's
QUALITY = 80

IMAGE = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".tif", ".tiff"}
VIDEO = {".mp4", ".mov", ".m4v", ".avi", ".webm"}
DOCUMENT = {".pdf", ".doc", ".docx", ".ppt", ".pptx", ".xls", ".xlsx"}

# a post counts when its link is on one of these: host suffix -> the platform's name
PLATFORMS = {"x.com": "X", "twitter.com": "X", "linkedin.com": "LinkedIn", "bsky.app": "Bluesky",
             "facebook.com": "Facebook", "instagram.com": "Instagram", "tiktok.com": "TikTok",
             "youtube.com": "YouTube", "youtu.be": "YouTube"}
# the only host a post's picture is taken from: its addresses are stable (LinkedIn's expire)
POST_IMAGE_HOST = "pbs.twimg.com"


def rendition_names(key):
    return [f"{key}-w{w}.jpg" for w in WIDTHS]


def _natural(s):
    return [int(p) if p.isdigit() else p.lower() for p in re.split(r"(\d+)", s)]


def _kind(name):
    ext = Path(name).suffix.lower()
    return "image" if ext in IMAGE else "video" if ext in VIDEO else "document" if ext in DOCUMENT else None


def _sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def read_folder(root):
    """Every file under root: [{rel, path, bytes, sha256}], folders in natural order."""
    root = Path(root)
    if not root.is_dir():
        raise SystemExit(f"FATAL: {root} is not a folder (the unzipped 'AA visuals')")
    files = [p for p in root.rglob("*") if p.is_file()]
    files.sort(key=lambda p: [_natural(x) for x in p.relative_to(root).parts])
    return [{"rel": p.relative_to(root).as_posix(), "path": p, "bytes": p.stat().st_size,
             "sha256": _sha256(p)} for p in files]


def _sheet(wb, name):
    """A sheet as a list of dicts keyed by its header row."""
    if name not in wb.sheetnames:
        raise SystemExit(f"FATAL: the catalogue has no sheet {name!r} (it has: {', '.join(wb.sheetnames)})")
    rows = list(wb[name].iter_rows(values_only=True))
    head = [str(h).strip() if h is not None else "" for h in rows[0]]
    return [dict(zip(head, r)) for r in rows[1:] if any(c not in (None, "") for c in r)]


def _s(v):
    return "" if v is None else str(v).strip()


def _platform(url):
    host = urllib.parse.urlsplit(url).netloc.lower().removeprefix("www.")
    return next((p for h, p in PLATFORMS.items() if host == h or host.endswith("." + h)), None)


def _date_from_id(url):
    """The day a post was made, read from its own identifier (UTC): X and LinkedIn ids start
    with the creation time in milliseconds, a Bluesky record key is its creation time."""
    try:
        if m := re.search(r"(?:x|twitter)\.com/[^/]+/status/(\d{15,20})", url):
            ms = (int(m.group(1)) >> 22) + 1288834974657
        elif m := re.search(r"linkedin\.com/.*?(?:activity|ugcPost)[-:](\d{19})", url):
            ms = int(m.group(1)) >> 22
        elif m := re.search(r"bsky\.app/profile/[^/]+/post/([2-7a-z]{13})\b", url):
            n = 0
            for ch in m.group(1):
                n = n * 32 + "234567abcdefghijklmnopqrstuvwxyz".index(ch)
            ms = (n >> 10) // 1000
        else:
            return ""
        d = datetime.fromtimestamp(ms / 1000, UTC).date()
        return d.isoformat() if date(2006, 1, 1) <= d <= date.today() else ""
    except (ValueError, OverflowError, OSError):
        return ""


def _post_image(url):
    """A post's picture, small rendition, when the catalogue has a stable address for it."""
    u = urllib.parse.urlsplit(_s(url))
    if u.scheme != "https" or u.netloc.lower() != POST_IMAGE_HOST:
        return ""
    q = dict(urllib.parse.parse_qsl(u.query))
    q["name"] = "small"
    return urllib.parse.urlunsplit((u.scheme, u.netloc, u.path, urllib.parse.urlencode(q), ""))


def read_posts(rows):
    """The catalogue's posts sheet -> (posts, left_out). Only the post itself is kept."""
    posts, left = [], {"no exact post link": 0, "not on a public platform": 0, "listed twice": 0}
    seen = set()
    for r in rows:
        url = _s(r.get("Exact individual post URL"))
        if not url.lower().startswith("https://"):
            left["no exact post link"] += 1
            continue
        platform = _platform(url)
        if not platform:
            left["not on a public platform"] += 1
            continue
        if url in seen:
            left["listed twice"] += 1
            continue
        seen.add(url)
        given = re.match(r"\d{4}-\d{2}-\d{2}", _s(r.get("Date")))
        post = {"platform": platform, "account": _s(r.get("Account")),
                "date": given.group(0) if given else _date_from_id(url),
                "country": _s(r.get("Country / region")), "theme": _s(r.get("Hazard / theme")),
                "asset": _s(r.get("Asset type")),
                "text": _s(r.get("Original post text / verified excerpt")), "url": url,
                "image": _post_image(r.get("Direct original media URL (when recoverable)"))}
        posts.append({k: v for k, v in post.items() if v})
    posts.sort(key=lambda p: (p.get("date", ""), p["url"]), reverse=True)
    return posts, {k: n for k, n in left.items() if n}


def build(root, files):
    """(manifest, notes) from the files of the folder."""
    import openpyxl

    at_root = [f for f in files if "/" not in f["rel"]]
    sheets = sorted((f for f in at_root if f["rel"].startswith(CATALOGUE_PREFIX) and f["rel"].endswith(".xlsx")),
                    key=lambda f: _natural(f["rel"]))
    if not sheets:
        raise SystemExit(f"FATAL: no workbook named {CATALOGUE_PREFIX!r}….xlsx at the root of the folder: "
                         "the posts come from it")
    cat = sheets[-1]                         # the highest version, when several are kept
    wb = openpyxl.load_workbook(cat["path"], read_only=True, data_only=True)
    posts, posts_left = read_posts(_sheet(wb, POSTS_SHEET))

    groups, skipped = {}, []
    for f in files:
        *path, name = f["rel"].split("/")
        kind = _kind(name)
        if name.startswith(".") or kind is None or (not path and Path(name).suffix.lower() in (".xls", ".xlsx")):
            skipped.append(f["rel"])
            continue
        groups.setdefault(tuple(path), []).append(
            {"key": f["sha256"][:16], "name": name, "kind": kind, "bytes": f["bytes"]})
    empty = sorted(p.relative_to(root).as_posix() for p in Path(root).rglob("*")
                   if p.is_dir() and not any(p.iterdir()))
    manifest = {"synced": date.today().isoformat(),
                "folder": {"title": Path(root).name, **SHAREPOINT}, "catalogue": Path(cat["rel"]).stem,
                "groups": [{"path": list(p), "items": items} for p, items in groups.items()],
                "posts": posts}
    return manifest, {"skipped": skipped, "empty": empty, "posts_left": posts_left}


def renditions(path):
    """The photo at path as {width: JPEG bytes}, upright, never wider than the original."""
    from PIL import Image, ImageOps

    with Image.open(path) as im:
        im.draft("RGB", (max(WIDTHS), max(WIDTHS)))   # a JPEG decodes at the scale needed, not full size
        icc = im.info.get("icc_profile") if im.mode == "RGB" else None
        im = ImageOps.exif_transpose(im)
        if im.mode != "RGB":
            im = im.convert("RGB")
        out = {}
        for w in WIDTHS:
            r = im.copy()
            r.thumbnail((w, 1 << 16), Image.LANCZOS)
            buf = io.BytesIO()
            r.save(buf, "JPEG", quality=QUALITY, optimize=True, progressive=True, icc_profile=icc)
            out[w] = buf.getvalue()
        return out


def _blobs(cc, prefix):
    """{name: size} of the blobs under prefix."""
    return {b.name: b.size for b in cc.list_blobs(name_starts_with=prefix + "/")}


def archive(files, manifest, dry_run):
    """Send what the blob lacks: the originals, the index, the renditions. Says what it sent."""
    import ocha_stratus as stratus
    from azure.storage.blob import ContentSettings

    cc = stratus.get_container_client(container_name=CONTAINER, stage="dev", write=not dry_run)
    have, web = _blobs(cc, ORIGINALS), _blobs(cc, WEB)
    send = [f for f in files if have.get(f"{ORIGINALS}/{f['rel']}") != f["bytes"]]
    photos = {i["key"]: i for g in manifest["groups"] for i in g["items"] if i["kind"] == "image"}
    by_key = {f["sha256"][:16]: f for f in files}
    make = [k for k in photos if any(f"{WEB}/{n}" not in web for n in rendition_names(k))]
    print(f"blob ({CONTAINER}, dev): {len(send)} original(s) to send "
          f"({sum(f['bytes'] for f in send) / 1e9:.2f} GB), {len(files) - len(send)} already there · "
          f"{len(make)} photo(s) to render, {len(photos) - len(make)} already rendered")
    if dry_run:
        return
    for k, f in enumerate(send, 1):
        ctype = mimetypes.guess_type(f["rel"])[0] or "application/octet-stream"
        with open(f["path"], "rb") as fh:
            cc.upload_blob(f"{ORIGINALS}/{f['rel']}", fh, overwrite=True, max_concurrency=4,
                           content_settings=ContentSettings(content_type=ctype))
        print(f"  original {k}/{len(send)}: {f['rel']}", flush=True)
    index = [{"path": f["rel"], "bytes": f["bytes"], "sha256": f["sha256"]} for f in files]
    stratus.upload_blob_data(json.dumps({"synced": manifest["synced"], "sharepoint": SHAREPOINT,
                                         "files": index}, ensure_ascii=False, indent=1).encode(),
                             INDEX, stage="dev", container_name=CONTAINER, content_type="application/json")
    total = 0
    for k, key in enumerate(make, 1):
        for w, data in renditions(by_key[key]["path"]).items():
            cc.upload_blob(f"{WEB}/{key}-w{w}.jpg", data, overwrite=True,
                           content_settings=ContentSettings(content_type="image/jpeg",
                                                            cache_control="public, max-age=31536000"))
            total += len(data)
        if k % 25 == 0 or k == len(make):
            print(f"  rendered {k}/{len(make)} ({total / 1e6:.0f} MB so far)", flush=True)


def fetch(out):
    """The renditions the manifest lists, from the blob into out (and nothing else there)."""
    import ocha_stratus as stratus
    from azure.core.exceptions import ResourceNotFoundError

    m = json.loads(MANIFEST.read_text()) if MANIFEST.exists() else {}
    want = sorted({n for g in m.get("groups", []) for i in g["items"] if i["kind"] == "image"
                   for n in rendition_names(i["key"])})  # a photo kept in two folders: one pair
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    for p in out.iterdir():                  # a cached folder: what the manifest dropped goes
        if p.name not in want:
            p.unlink()
    cc = stratus.get_container_client(container_name=CONTAINER, stage="dev")
    got = 0
    for n in want:
        p = out / n
        if p.exists() and p.stat().st_size:  # the name is the photo's content: never stale
            continue
        try:
            p.write_bytes(cc.download_blob(f"{WEB}/{n}").readall())
        except ResourceNotFoundError as e:
            raise SystemExit(f"FATAL: rendition {WEB}/{n} is not on the blob: "
                             "run sync_media.py --from on the folder") from e
        got += 1
    print(f"media: {len(want)} rendition(s) in {out} ({got} fetched, {len(want) - got} cached)")


def dumps(m):
    """The manifest as JSON with one file and one post per line, so a sync reads as a diff."""
    def j(v):
        return json.dumps(v, ensure_ascii=False)
    out = ["{", f' "synced": {j(m["synced"])},', f' "folder": {j(m["folder"])},',
           f' "catalogue": {j(m["catalogue"])},', ' "groups": [']
    for gi, g in enumerate(m["groups"]):
        out.append(f'  {{"path": {j(g["path"])}, "items": [')
        out += [f"   {j(it)}{',' if k < len(g['items']) - 1 else ''}" for k, it in enumerate(g["items"])]
        out.append("  ]}" + ("," if gi < len(m["groups"]) - 1 else ""))
    out += [" ],", ' "posts": [']
    out += [f"  {j(p)}{',' if k < len(m['posts']) - 1 else ''}" for k, p in enumerate(m["posts"])]
    out += [" ]", "}"]
    return "\n".join(out) + "\n"


def _changes(old, new):
    """What a sync changes, in words: files (by folder and name) and posts that came and went."""
    def files(m):
        return {"/".join(g["path"] + [i["name"]]) for g in m.get("groups", []) for i in g["items"]}
    def posts(m):
        return {p["url"]: f'{p["platform"]} {p.get("date", "")} {p.get("account", "")}' for p in m.get("posts", [])}
    lines = []
    for what, a, b in (("file", {f: f for f in files(old)}, {f: f for f in files(new)}),
                       ("post", posts(old), posts(new))):
        for sign, ids, src in (("+", b.keys() - a.keys(), b), ("-", a.keys() - b.keys(), a)):
            names = sorted(src[i] for i in ids)
            if names:
                lines.append(f"  {sign}{len(names)} {what}(s): " + "; ".join(names[:8])
                             + (f"; … {len(names) - 8} more" if len(names) > 8 else ""))
    return lines


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--from", dest="root", metavar="DIR", help="the unzipped 'AA visuals' folder")
    src.add_argument("--fetch", metavar="DIR", help="fetch the manifest's renditions into DIR (publish)")
    ap.add_argument("--dry-run", action="store_true", help="say what would change, write and send nothing")
    a = ap.parse_args()
    if a.fetch:
        return fetch(a.fetch)

    files = read_folder(a.root)
    manifest, notes = build(a.root, files)
    n = {k: sum(1 for g in manifest["groups"] for i in g["items"] if i["kind"] == k)
         for k in ("image", "video", "document")}
    size = sum(f["bytes"] for f in files)
    print(f'{manifest["folder"]["title"]}: {len(files)} files ({size / 1e9:.2f} GB) · listed: {n["image"]} photos, '
          f'{n["video"]} videos, {n["document"]} documents in {len(manifest["groups"])} folders · '
          f'{len(manifest["posts"])} posts from {manifest["catalogue"]!r}')
    for k, v in notes["posts_left"].items():
        print(f"  left out: {v} catalogue row(s) — {k}")
    if notes["skipped"]:
        print(f'  not listed (archived only): {len(notes["skipped"])} file(s) — ' + "; ".join(notes["skipped"]))
    if notes["empty"]:
        print("  empty folder(s): " + "; ".join(notes["empty"]))
    dupes = len(files) - len({f["sha256"][:16] for f in files})
    if dupes:
        print(f"  {dupes} file(s) are copies of another file of the folder (same content, same key)")
    old = json.loads(MANIFEST.read_text()) if MANIFEST.exists() else {}
    changes = _changes(old, manifest)
    print("\n".join(changes) if changes else "  no file or post added or removed")
    archive(files, manifest, a.dry_run)
    if a.dry_run:
        print("dry run: nothing written, nothing sent")
        return
    MANIFEST.write_text(dumps(manifest))
    print(f"wrote {MANIFEST.relative_to(Path(__file__).parents[1])} — read the diff before committing it")


if __name__ == "__main__":
    sys.exit(main())
