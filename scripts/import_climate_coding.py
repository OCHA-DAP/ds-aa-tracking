"""Turn a climate-adaptation coding workbook into a data-entry file for aa.cerf_climate_coding.

The workbook is the coder's sheet: two header rows (area titles over the indicator
questions), then one row per CERF project with the project's facts in A:K and the answers
in L:AS (Y/N per indicator; five "other relevant activities" columns hold free text). The
columns are read by position and every header is checked against the expected wording
first, so a sheet laid out differently stops the script instead of shifting answers onto
the wrong indicator.

Cells are normalised once: y/n/yes/no/TRUE/FALSE -> Y/N; "Y (?)", "N?" -> Y/N marked
uncertain; an Excel error value (#VALUE!) -> no answer; any other text typed into a Y/N
cell (a remark) -> no answer, the remark kept as the note. A cell comment becomes the note
of that answer. Every indicator of every project is written, a blank one as NULL, so
re-importing a corrected workbook also clears the answers the coder took back.

Usage (aa.cerf_climate_coding must exist first: the nightly job's one-off
`--ensure-schema` run creates it; a file applied before that fails, and so do the files
after it that night):
    uv run python scripts/import_climate_coding.py WORKBOOK --coder NAME [--sheet "Sample 50"] --out FILE.json
    uv run python scripts/apply_entries.py --upload FILE.json      # the nightly job applies it

The file goes to the PRIVATE blob only, never to git (this repo is public)."""

import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from page_climate import KEYS  # noqa: E402

# column -> (indicator key, start of the header wording in row 2); a None key is the free
# text of the "other" indicator just before it
COLUMNS = {
    "L": ("climate_risk_mentioned", "Are current climate risk or future impacts mentioned"),
    "M": ("water_sanitation", "activities help increase or extend"),
    "N": ("water_stress", "activities reduce water stress"),
    "O": ("water_efficiency", "activities increase water efficiency"),
    "P": ("water_relocation", "activiites support planned relocaiton"),
    "Q": ("water_other", "Project has other relevant activities"),
    "R": (None, "Other relevant activities"),
    "S": ("food_extension", "support for extension services"),
    "T": ("food_adaptation_tech", "support for agriculture that uses adaptation technology"),
    "U": ("food_access", "activities increase proportion of population with equitable access"),
    "V": ("food_yield", "activities increase yield"),
    "W": ("food_other", "Project supports other climate relevant activity"),
    "X": (None, "Other relevant activities"),
    "Y": ("health_infectious_disease", "Decrease incidence of climate related infectious disease"),
    "Z": ("health_mortality_morbidity", "Reduced mortality or morbidity"),
    "AA": ("health_mental_health", "Increased proportion of population"),
    "AB": ("health_services_capacity", "Support for health services to remain at full"),
    "AC": ("health_facilities_resilient", "Support for health facilities to be resilient"),
    "AD": ("health_other", "Other relevant activities to heath"),
    "AE": (None, "Other relevant activities"),
    "AF": ("ecosystems_activities", "Project includes relevant activities"),
    "AG": (None, "Other relevant activities"),
    "AH": ("infra_relocation", "infrustrucuture or settlement relocated"),
    "AI": ("infra_settlement_upgrade", "settlement upgrade that includes adaptation"),
    "AJ": ("infra_other", "Other relevant activities to heath"),   # the sheet's label (sic)
    "AK": (None, "Other relevant activities"),
    "AL": ("poverty_social_protection_access", "increased access to social protection"),
    "AM": ("poverty_social_protection_climate", "social protection services consider climate risk"),
    "AN": ("poverty_reduced", "Reduced population living in poverty"),
    "AO": ("ews_multi_hazard", "help establish multi-hazards EWS"),
    "AP": ("ews_warning_access", "increase number of people that have access to warnings"),
    "AQ": ("ews_monitoring_forecasts", "support for hazards monitoring and forecasts"),
    "AR": ("ews_evacuation", "increased population evaculated"),
    "AS": ("ews_climate_information", "support for climate information services"),
}
# only guards against a shifted column A: CERF codes have taken several shapes since 2006
# (09-IOM-026, 10-CEF-010-A, 14-RR-CEF-003, CERF-AFG-25-RR-CEF-35353)
CODE_RE = re.compile(r"^(\d\d-|CERF-)[A-Z0-9()-]+$")


def _norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def answer(v):
    """A Y/N cell -> (value, uncertain, remark)."""
    if isinstance(v, bool):
        return ("Y" if v else "N"), False, None
    s = "" if v is None else str(v).strip()
    if s.upper() in ("Y", "N", "YES", "NO"):
        return s[0].upper(), False, None
    if m := re.fullmatch(r"([YyNn])\s*\(?\?\)?", s):    # Y (?), N?, …: unsure
        return m.group(1).upper(), True, None
    if re.fullmatch(r"#[A-Z/0]+[!?]?", s):       # an Excel error value pasted in: no answer
        return None, False, None
    return None, False, (s or None)


def other_text(v):
    """A free-text 'other activities' cell; N, N/A, a dash or 'none' mean none."""
    s = "" if v is None else str(v).strip()
    return None if s.upper() in ("", "N", "NO", "N/A", "NA", "-", "NONE") else s


def read(path, sheet, coder):
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]     # cached values, not formulas
    if _norm(ws["A2"].value) != "project code":
        raise SystemExit(f"{sheet}!A2 is {ws['A2'].value!r}, expected 'Project Code' (two header rows)")
    for col, (key, head) in COLUMNS.items():
        got = _norm(ws[f"{col}2"].value)
        # a free-text column is matched whole, so it cannot pass for the Y/N column before it
        if (got != _norm(head)) if key is None else not got.startswith(_norm(head)):
            raise SystemExit(f"{sheet}!{col}2 is {ws[f'{col}2'].value!r}, expected {head!r}")
    keys = [k for k, _ in COLUMNS.values() if k]
    if set(keys) != KEYS:
        raise SystemExit(f"columns and page_climate.INDICATORS disagree: {sorted(set(keys) ^ KEYS)}")

    source = f"{Path(path).name} [{sheet}]"
    rows, blank, odd, seen = [], [], [], {}
    for r in range(3, ws.max_row + 1):
        code = str(ws[f"A{r}"].value or "").strip()
        if not code:
            continue
        if not CODE_RE.match(code):
            raise SystemExit(f"{sheet}!A{r}: {code!r} is not a CERF project code")
        if code in seen:
            raise SystemExit(f"{sheet}!A{r}: {code} is already on row {seen[code]}")
        seen[code] = r
        got, last = [], None
        for col, (key, _) in COLUMNS.items():
            cell = ws[f"{col}{r}"]
            comment = cell.comment.text.split("Comment:")[-1].strip() if cell.comment else None
            if key is None:                      # the free text of the 'other' indicator before it
                last["note"] = "; ".join(x for x in (last["note"], other_text(cell.value), comment) if x) or None
                continue
            value, unc, remark = answer(cell.value)
            if remark:
                odd.append(f"{code} {key}: {remark!r}")
            note = "; ".join(x for x in (remark, comment) if x) or None
            last = {"project_code": code, "indicator": key, "coded_by": coder, "value": value,
                    "uncertain": unc, "note": note, "source": source}
            got.append(last)
        # every indicator is written, blank ones as NULL: a re-import then clears an answer
        # the coder took back, instead of leaving the old one in place
        if not any(g["value"] for g in got):
            blank.append(code)
        rows += got
    return rows, blank, odd


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("workbook")
    ap.add_argument("--coder", required=True, help="who coded the sheet (aa.cerf_climate_coding.coded_by)")
    ap.add_argument("--sheet", default="Sample 50")
    ap.add_argument("--out", required=True, help="the entries file to write (JSON)")
    a = ap.parse_args()
    rows, blank, odd = read(a.workbook, a.sheet, a.coder)
    projects = {r["project_code"] for r in rows}
    payload = {"entered_by": f"{a.coder} (coding workbook {Path(a.workbook).name})",
               "note": f"climate-adaptation coding, sheet {a.sheet!r}: {len(projects)} projects",
               "rows": [{"table": "cerf_climate_coding", "row": r} for r in rows]}
    Path(a.out).write_text(json.dumps(payload, ensure_ascii=False, indent=1))
    print(f"{a.out}: {len(rows)} rows, {len(projects)} projects")
    if blank:
        print(f"  no answers at all (written as blanks, listed on the page as not coded): {', '.join(blank)}")
    for o in odd:
        print(f"  not a Y/N answer, kept as the note: {o}")


if __name__ == "__main__":
    main()
